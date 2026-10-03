# -*- coding: utf-8 -*-
"""Шаблон сборки русского доклада .docx по домовому образцу (см. skill tex-to-doc).

Запуск: uv run --with python-docx build_doc.py
Затем:  soffice --headless --convert-to doc <OUT>.docx --outdir .

ЗАПОЛНИТЬ: CONFIG, шапку (УДК/название/авторы/организация/e-mail), блоки B, REFS.
"""
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============================ CONFIG ============================
FIG = "/abs/path/to/figure_ru.png"     # PNG рисунка (если есть)
OUT = "/abs/path/to/report.docx"

FONT = "Times New Roman"
BODY = 14    # кегль основного текста, пт
SMALL = 12   # кегль подписей и таблиц, пт

TAG_RE = re.compile(r"\{(i|b|sub|sup|isub|isup)\}(.*?)\{/\1\}", re.S)


def parse(text):
    parts, pos = [], 0
    for m in TAG_RE.finditer(text):
        if m.start() > pos:
            parts.append(("t", text[pos:m.start()]))
        parts.append((m.group(1), m.group(2)))
        pos = m.end()
    if pos < len(text):
        parts.append(("t", text[pos:]))
    return parts


def style_run(r, size=BODY, bold=None):
    r.font.name = FONT
    r.font.size = Pt(size)
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    if bold:
        r.bold = True


def add_rich(p, parts, size=BODY):
    for kind, text in parts:
        r = p.add_run(text)
        style_run(r, size=size)
        if kind == "i":
            r.italic = True
        elif kind == "b":
            r.bold = True
        elif kind == "sub":
            r.font.subscript = True
        elif kind == "sup":
            r.font.superscript = True
        elif kind == "isub":
            r.italic = True
            r.font.subscript = True
        elif kind == "isup":
            r.italic = True
            r.font.superscript = True
    return p


def suppress_hyphen(p):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:suppressAutoHyphens")
    el.set(qn("w:val"), "true")
    pPr.insert_element_before(
        el,
        "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
        "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
        "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
        "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
        "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(1), size=BODY,
         no_hyph=False, before=0, after=0, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = 1.0
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = indent
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    if keep_next:
        pf.keep_with_next = True
    add_rich(p, parse(text), size=size)
    if no_hyph:
        suppress_hyphen(p)
    return p


def equation(doc, markup, number, two_lines=None):
    """Выключная формула текстовыми run'ами (НЕ OMML — ломается при .doc).
    Табуляция: центр 8,5 см, правая 17 см (правая = ширина листа − поля)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.0
    pf.space_before = Pt(5)
    pf.space_after = Pt(5)
    pf.first_line_indent = Cm(0)
    pf.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
    suppress_hyphen(p)
    p.add_run("\t")
    add_rich(p, parse(markup))
    if two_lines:                     # вторая строка длинной формулы
        p.add_run("\n\t")
        add_rich(p, parse(two_lines))
    p.add_run("\t(" + number + ")")
    return p


def set_cell(cell, text, align, size=SMALL):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    add_rich(p, parse(text), size=size)


def add_table(doc, caption, headers, rows, widths_cm, first_col_left=True):
    para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Cm(0),
         size=SMALL, no_hyph=True, before=8, after=4, keep_next=True)
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    for j, h in enumerate(headers):
        c = table.cell(0, j)
        c.width = Cm(widths_cm[j])
        set_cell(c, h, WD_ALIGN_PARAGRAPH.CENTER)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = table.cell(i + 1, j)
            c.width = Cm(widths_cm[j])
            al = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and first_col_left) else WD_ALIGN_PARAGRAPH.CENTER
            set_cell(c, val, al)
    # держать таблицу целиком на одной странице (все строки, кроме последней)
    for i, row in enumerate(table.rows):
        keep = i < len(table.rows) - 1
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = keep
    para(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0), size=2, after=2)


doc = Document()

# ---------- страница ----------
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(2)
sec.right_margin = Cm(2)
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(3.2)
sec.header_distance = Cm(1.1)
sec.different_first_page_header_footer = True   # на титуле номера нет

hp = sec.header.paragraphs[0]
hp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
hp.paragraph_format.first_line_indent = Cm(0)
run = hp.add_run()
style_run(run, size=SMALL)
fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
itxt = OxmlElement("w:instrText"); itxt.set(qn("xml:space"), "preserve"); itxt.text = "PAGE \\* MERGEFORMAT"
fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
run._r.append(fld1); run._r.append(itxt); run._r.append(fld2)

# ---------- шрифт по умолчанию ----------
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(BODY)
st.paragraph_format.line_spacing = 1.0
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(0)
rpr = st.element.get_or_add_rPr()
lang = OxmlElement("w:lang")
lang.set(qn("w:val"), "ru-RU")
rpr.append(lang)

settings = doc.settings.element
ah = OxmlElement("w:autoHyphenation"); ah.set(qn("w:val"), "true")
settings.append(ah)
hz = OxmlElement("w:hyphenationZone"); hz.set(qn("w:val"), "425")
settings.append(hz)

# ---------- шапка ----------
para(doc, "УДК 000.00", align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(0))
para(doc, "НАЗВАНИЕ ДОКЛАДА ЗАГЛАВНЫМИ БУКВАМИ",
     align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0), no_hyph=True, before=12,
     keep_next=True)
for r in doc.paragraphs[-1].runs:
    r.bold = True
para(doc, "И. О. Фамилия, И. О. Фамилия",
     align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0), before=6, keep_next=True)
para(doc, "Организация, город, страна",
     align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0), before=2, keep_next=True)
for r in doc.paragraphs[-1].runs:
    r.italic = True
para(doc, "{b}E-mail:{/b} author@example.org",
     align=WD_ALIGN_PARAGRAPH.LEFT, indent=Cm(0), before=2)

# ---------- текст ----------
B = []

B.append(("p", "{b}Введение. {/b}Первый абзац раздела с врезанным полужирным "
              "заголовком. Разделы без нумерации; ссылки на источники в квадратных "
              "скобках [1]. Мини-разметка: {i}курсив{/i}, {b}полужирный{/b}, "
              "{sub}нижний{/sub}, {sup}верхний{/sup} индекс."))
# B.append(("eq", "формула", "1", None))                 # выключная формула (1)
# B.append(("eq", "первая строка", "1", "вторая строка"))  # двухстрочная
# B.append(("table", "Таблица 1. Подпись 12 пт",
#           ["Столбец 1", "Столбец 2"],
#           [["значение", "значение"]],
#           [6.0, 6.0]))                                  # ширины столбцов, см
# B.append(("fig", FIG, 11.0, "Рис. 1. Подрисуночная подпись 12 пт"))

# ---------- список источников ----------
para(doc, "{b}Список использованных источников{/b}", align=WD_ALIGN_PARAGRAPH.JUSTIFY,
     indent=Cm(1), before=10, keep_next=True)

# ГОСТ 7.05-2008, нумерация по порядку цитирования, формат:
# «Фамилия, И. О. Название / И. О. Фамилия [et al.] // Журнал. – Год.»
REFS = [
    "1. Фамилия, И. О. Название / И. О. Фамилия, И. О. Фамилия [et al.] // Журнал. – 2024.",
]

for ref in REFS:
    para(doc, ref, after=4)

doc.save(OUT)
print("saved", OUT)
